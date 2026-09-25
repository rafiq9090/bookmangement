import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def build_clean_report():
    doc = docx.Document()

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    primary_color = RGBColor(26, 54, 93)   # Navy Blue
    secondary_color = RGBColor(74, 85, 104) # Muted Slate

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def add_clean_heading(text, level):
        h = doc.add_heading(text, level=level)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        run = h.runs[0]
        run.font.name = 'Calibri'
        if level == 1:
            run.font.size = Pt(15)
            run.font.color.rgb = primary_color
            run.bold = True
        elif level == 2:
            run.font.size = Pt(12)
            run.font.color.rgb = primary_color
            run.bold = True
        return h

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run('Old & Rare Book Multi-Vendor Marketplace')
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = primary_color

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run('Software Engineering Project Report')
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = secondary_color

    # Meta Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ('Project Lead & Full-Stack Architect', 'Md. Rafiqul Islam'),
        ('Frontend & UI/UX Developer', '[Team Member 2 Name]'),
        ('Database & QA Engineer', '[Team Member 3 Name]'),
        ('Technology Stack', 'Django 5.1, PostgreSQL 16, Redis 7, Celery, Tailwind CSS, Docker')
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.width = Inches(2.6)
        cell_v.width = Inches(3.9)
        set_cell_background(cell_k, 'F1F5F9')
        set_cell_background(cell_v, 'FFFFFF')
        
        pk = cell_k.paragraphs[0]
        rk = pk.add_run(k)
        rk.bold = True
        rk.font.name = 'Calibri'
        rk.font.size = Pt(10)

        pv = cell_v.paragraphs[0]
        rv = pv.add_run(v)
        rv.font.name = 'Calibri'
        rv.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 1. Executive Summary
    add_clean_heading('1. Executive Summary', 1)
    p = doc.add_paragraph(
        'Traditional e-commerce platforms assume products are brand-new and identical. '
        'In contrast, antique, out-of-print, and pre-owned books are each unique—differing in wear, paper condition, edition, annotations, and binding. '
        'The Old & Rare Book Marketplace is a dedicated multi-vendor web platform engineered specifically for this domain. '
        'It introduces physical condition grading, direct buyer-seller inquiries, zero double-selling guarantees via atomic locks, '
        'automatic multi-vendor order splitting, and escrow-backed payments.'
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)

    # 2. Key Objectives & Features
    add_clean_heading('2. Key Features & Objectives', 1)
    features = [
        ('Accurate Physical Grading', 'Sellers classify listings using 6 standard condition grades (As New, Fine, Very Good, Good, Fair, Poor) supported by real inspection photos and condition notes.'),
        ('Buyer-Seller Inquiry Chat', 'Prospective buyers can message sellers directly to request close-up photos or negotiate pricing before unlocking checkout.'),
        ('Zero Double-Selling (Atomic Locks)', '15-minute Redis reservation locks guarantee that unique single-copy books cannot be purchased concurrently by multiple buyers.'),
        ('Multi-Vendor Order Splitting', 'Buyers can order from multiple sellers in one transaction; the system automatically splits the order into separate shipments with distinct tracking and payouts.'),
        ('Escrow Payment Security', 'Buyer payments are held in escrow and released to the seller ledger only upon confirmed delivery, protecting both parties against fraud.')
    ]
    for title, desc in features:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(f'{title}: ')
        r1.bold = True
        r1.font.name = 'Calibri'
        r2 = bp.add_run(desc)
        r2.font.name = 'Calibri'

    # 3. Technology Stack
    add_clean_heading('3. Technology Stack', 1)
    tech_table = doc.add_table(rows=7, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_rows = [
        ('Component', 'Technology', 'Purpose in Project'),
        ('Backend Framework', 'Django 5.1 & Django REST Framework', 'Modular MVC architecture, robust ORM, authentication, and REST APIs'),
        ('Database', 'PostgreSQL 16', 'ACID transactions, relational constraints, and JSONB event logs'),
        ('Caching & Locks', 'Redis 7', '15-minute checkout reservation locks and session caching'),
        ('Asynchronous Tasks', 'Celery + Redis Broker', 'WebP image compression, expired lock cleanup, and email notifications'),
        ('Frontend & UI', 'Tailwind CSS + HTML5 / JS', 'Modern, responsive interface with interactive modal showcases'),
        ('Containerization', 'Docker & Docker Compose', 'Containerized web application, PostgreSQL database, and Redis cache')
    ]
    for i, row_data in enumerate(tech_rows):
        row = tech_table.rows[i]
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            if i == 0:
                set_cell_background(cell, '1A365D')
                p = cell.paragraphs[0]
                r = p.add_run(val)
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, 'F8FAFC' if i % 2 == 1 else 'FFFFFF')
                p = cell.paragraphs[0]
                r = p.add_run(val)
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 4. Core System Modules
    add_clean_heading('4. Core System Modules', 1)
    modules = [
        ('Accounts (apps.accounts)', 'Custom user model with role-based access (Buyer, Seller, Admin), seller KYC verification, and delivery address management.'),
        ('Book Catalog (apps.books)', 'Master catalog for canonical book metadata (ISBN, title, author, category) kept separate from individual seller copies to prevent duplicate records.'),
        ('Listings (apps.listings)', 'Seller inventory linked to master books, containing condition grades, seller notes, price, and real-copy inspection photos.'),
        ('Orders & Fulfillment (apps.orders)', 'Multi-item cart, concurrency-safe checkout via atomic locks, and automated splitting into per-seller shipments.'),
        ('Messaging & Inquiries (apps.messaging)', 'Inquiry threads attached to specific book copies for condition verification and seller checkout approval.'),
        ('Payments & Escrow (apps.payments)', 'Payment gateway integration with an escrow hold mechanism that releases earnings to seller ledgers only upon verified delivery.'),
        ('Shipping & Logistics (apps.shipping)', 'Pluggable courier integrations (Pathao, Steadfast) with chronological delivery status updates.')
    ]
    for mod_title, mod_desc in modules:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(f'{mod_title}: ')
        r1.bold = True
        r1.font.name = 'Calibri'
        r2 = bp.add_run(mod_desc)
        r2.font.name = 'Calibri'

    # 5. Work Distribution
    add_clean_heading('5. Work Distribution & Contributions', 1)
    work_table = doc.add_table(rows=4, cols=3)
    work_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    work_rows = [
        ('Team Member', 'Role', 'Key Responsibilities'),
        ('Md. Rafiqul Islam\n(Project Lead)', 'Lead Full-Stack Developer & System Architect', 'System architecture, Django app structure, Books/Listings/Orders logic, Redis locks, multi-vendor shipment splitting, and Escrow payments.'),
        ('[Team Member 2 Name]', 'Frontend & UI/UX Developer', 'Tailwind CSS responsive design, interactive book showcase, photo switchers, cart/checkout interfaces, and mobile optimization.'),
        ('[Team Member 3 Name]', 'Database, QA & DevOps Engineer', 'PostgreSQL database modeling, Docker Compose environment setup, automated API testing, test data fixtures, and report documentation.')
    ]
    for i, row_data in enumerate(work_rows):
        row = work_table.rows[i]
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            if i == 0:
                set_cell_background(cell, '1A365D')
                p = cell.paragraphs[0]
                r = p.add_run(val)
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, 'F8FAFC' if i % 2 == 1 else 'FFFFFF')
                p = cell.paragraphs[0]
                r = p.add_run(val)
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 6. Future Roadmap
    add_clean_heading('6. Future Roadmap', 1)
    roadmap = [
        ('AI Condition Assessment', 'Computer vision integration to inspect book photos, automatically detecting wear, spine condition, and water damage.'),
        ('Mobile Application (Flutter / React Native)', 'Mobile app with barcode scanner to let sellers list books instantly by scanning ISBNs.'),
        ('Real-Time Chat (WebSockets)', 'Instant buyer-seller messaging with typing indicators via Django Channels.'),
        ('Rare Book Auctions', 'Bidding and timer engine for high-value antique books, first editions, and manuscripts.')
    ]
    for r_title, r_desc in roadmap:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(f'{r_title}: ')
        r1.bold = True
        r1.font.name = 'Calibri'
        r2 = bp.add_run(r_desc)
        r2.font.name = 'Calibri'

    # 7. Conclusion
    add_clean_heading('7. Conclusion', 1)
    p_concl = doc.add_paragraph(
        'The Old & Rare Book Marketplace delivers a tailored solution for vintage and collectible book trading. '
        'By combining condition grading, pre-order condition inquiries, atomic concurrency locks, multi-vendor shipment splitting, and secure escrow accounting, '
        'the platform provides a reliable, fraud-resistant, and user-friendly experience for book collectors and independent sellers.'
    )
    p_concl.paragraph_format.line_spacing = 1.15

    output_path = 'Old_and_Rare_Book_Marketplace_Project_Report_Clean.docx'
    doc.save(output_path)
    print(f'Successfully generated {output_path}')

if __name__ == '__main__':
    build_clean_report()
